"""Import the canonical workbook to PostgreSQL, then index changed cars in Qdrant.

Run from the repository root: uv run python -m scripts.ingest_inventory
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session

from app.core.config import Settings, load_settings
from app.core.database import ensure_schema, make_engine
from app.models.car import CANONICAL_COLUMNS, Car, CarEmbeddingState
from app.services.embedding_service import EmbeddingError, EmbeddingService, text_hash
from app.services.vector_store import VectorStore, point_id
from app.services.vehicle_facts import extract_vehicle_facts

DEFAULT_WORKBOOK = (
    Path(__file__).resolve().parents[1] / "data" / "processed" / "cars_merged.xlsx"
)
REQUIRED = set(CANONICAL_COLUMNS)
FACT_COLUMNS = ("price_aed", "price_evidence", "mileage_km", "mileage_evidence")


@dataclass
class IngestionReport:
    source_rows: int = 0
    postgres_inserted: int = 0
    postgres_updated: int = 0
    postgres_canonical_rows: int = 0
    embeddings_requested: int = 0
    embeddings_skipped_unchanged: int = 0
    qdrant_upserted: int = 0
    qdrant_vector_count: int = 0
    failures: int = 0


def read_inventory(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        raise ValueError(f"Canonical workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "merged_inventory" not in workbook.sheetnames:
            raise ValueError("Workbook has no merged_inventory sheet")
        cells = workbook["merged_inventory"].iter_rows(values_only=True)
        headers = next(cells, None)
        if headers is None or len(headers) != len(set(headers)):
            raise ValueError("Missing or duplicate workbook headers")
        missing = REQUIRED - set(headers)
        if missing:
            raise ValueError(
                f"Missing required workbook columns: {', '.join(sorted(missing))}"
            )
        rows: list[dict[str, Any]] = []
        seen: set[str] = set()
        for excel_row, values in enumerate(cells, 2):
            if not any(value is not None for value in values):
                continue
            row = dict(zip(headers, values, strict=True))
            listing_id = row["listing_id"]
            if (
                not isinstance(listing_id, str)
                or not listing_id.strip()
                or listing_id in seen
            ):
                raise ValueError(
                    f"Missing or duplicate listing_id at Excel row {excel_row}"
                )
            if (
                not isinstance(row["embedding_text"], str)
                or not row["embedding_text"].strip()
            ):
                raise ValueError(f"Empty embedding_text at Excel row {excel_row}")
            if not isinstance(row["year"], int) or not 1900 <= row["year"] <= 2100:
                raise ValueError(f"Invalid year at Excel row {excel_row}")
            for field in REQUIRED - {
                "source_listing_id",
                "model_original",
                "trim_original",
            }:
                if row[field] is None:
                    raise ValueError(f"Null {field} at Excel row {excel_row}")
            seen.add(listing_id)
            rows.append({field: row[field] for field in REQUIRED})
        if not rows:
            raise ValueError("No canonical vehicles in merged_inventory")
        return rows
    finally:
        workbook.close()


def upsert_postgres(
    engine: Engine, rows: list[dict[str, Any]], report: IngestionReport
) -> None:
    ids = [row["listing_id"] for row in rows]
    enriched = [
        {**row, **extract_vehicle_facts(row["title"], row["description"])}
        for row in rows
    ]
    with engine.begin() as connection:
        prior = set(
            connection.scalars(select(Car.listing_id).where(Car.listing_id.in_(ids)))
        )
        for offset in range(0, len(enriched), 50):
            batch = enriched[offset : offset + 50]
            statement = insert(Car).values(batch)
            statement = statement.on_conflict_do_update(
                index_elements=[Car.listing_id],
                set_={
                    field: statement.excluded[field]
                    for field in REQUIRED.union(FACT_COLUMNS)
                    if field != "listing_id"
                },
            )
            connection.execute(statement)
        report.postgres_inserted = len(rows) - len(prior)
        report.postgres_updated = len(prior)
        report.postgres_canonical_rows = (
            connection.scalar(
                select(func.count()).select_from(Car).where(Car.listing_id.in_(ids))
            )
            or 0
        )


def _states(engine: Engine, ids: list[str]) -> dict[str, CarEmbeddingState]:
    with Session(engine) as session:
        return {
            row.listing_id: row
            for row in session.execute(
                select(CarEmbeddingState).where(CarEmbeddingState.listing_id.in_(ids))
            ).scalars()
        }


def index_vectors(
    engine: Engine,
    rows: list[dict[str, Any]],
    settings: Settings,
    report: IngestionReport,
) -> None:
    store = VectorStore(settings)
    ids = [row["listing_id"] for row in rows]
    states = _states(engine, ids)
    collection_exists = store.client.collection_exists(store.collection)
    existing = store.existing_points(ids) if collection_exists else {}
    pending: list[dict[str, Any]] = []
    reconciled: list[dict[str, Any]] = []
    for row in rows:
        row["embedding_hash"] = text_hash(row["embedding_text"])
        state = states.get(row["listing_id"])
        point = existing.get(point_id(row["listing_id"]))
        payload = point.payload if point else None
        current = (
            payload is not None
            and payload.get("embedding_hash") == row["embedding_hash"]
            and payload.get("embedding_model") == settings.nvidia_model
        )
        if current:
            report.embeddings_skipped_unchanged += 1
            if (
                state is None
                or state.embedding_hash != row["embedding_hash"]
                or state.model != settings.nvidia_model
            ):
                reconciled.append(row)
            metadata = {
                field: row[field]
                for field in ("listing_id", "year", "make", "model", "trim")
            }
            if any(payload.get(field) != value for field, value in metadata.items()):
                store.client.set_payload(
                    collection_name=store.collection,
                    payload=metadata,
                    points=[point_id(row["listing_id"])],
                    wait=True,
                )
        else:
            pending.append(row)

    if reconciled:
        config = store.client.get_collection(store.collection).config.params.vectors
        dimensions = (
            next(iter(config.values())).size
            if isinstance(config, dict)
            else config.size
        )
        with engine.begin() as connection:
            statement = insert(CarEmbeddingState).values(
                [
                    {
                        "listing_id": row["listing_id"],
                        "embedding_hash": row["embedding_hash"],
                        "model": settings.nvidia_model,
                        "dimensions": dimensions,
                    }
                    for row in reconciled
                ]
            )
            connection.execute(
                statement.on_conflict_do_update(
                    index_elements=[CarEmbeddingState.listing_id],
                    set_={
                        "embedding_hash": statement.excluded.embedding_hash,
                        "model": statement.excluded.model,
                        "dimensions": statement.excluded.dimensions,
                        "indexed_at": func.now(),
                    },
                )
            )

    if not pending:
        if collection_exists:
            dims = next(iter(states.values())).dimensions if states else None
            if dims:
                store.ensure_collection(dims)
        report.qdrant_vector_count = store.client.count(
            store.collection, exact=True
        ).count
        return

    service = EmbeddingService(settings)
    try:
        for offset in range(0, len(pending), settings.nvidia_batch_size):
            batch = pending[offset : offset + settings.nvidia_batch_size]
            vectors = service.embed_passages([row["embedding_text"] for row in batch])
            report.embeddings_requested += len(batch)
            dims = len(vectors[0])
            if any(len(vector) != dims for vector in vectors):
                raise ValueError("NVIDIA returned inconsistent vector dimensions")
            store.ensure_collection(dims)
            for chunk_start in range(0, len(batch), 32):
                chunk = batch[chunk_start : chunk_start + 32]
                store.upsert(
                    chunk,
                    vectors[chunk_start : chunk_start + 32],
                    settings.nvidia_model,
                )
                # Mark each chunk current only after Qdrant acknowledges its upsert.
                with engine.begin() as connection:
                    statement = insert(CarEmbeddingState).values(
                        [
                            {
                                "listing_id": row["listing_id"],
                                "embedding_hash": row["embedding_hash"],
                                "model": settings.nvidia_model,
                                "dimensions": dims,
                            }
                            for row in chunk
                        ]
                    )
                    connection.execute(
                        statement.on_conflict_do_update(
                            index_elements=[CarEmbeddingState.listing_id],
                            set_={
                                "embedding_hash": statement.excluded.embedding_hash,
                                "model": statement.excluded.model,
                                "dimensions": statement.excluded.dimensions,
                                "indexed_at": func.now(),
                            },
                        )
                    )
                report.qdrant_upserted += len(chunk)
    finally:
        service.close()
    report.qdrant_vector_count = store.client.count(store.collection, exact=True).count


def verify(
    engine: Engine,
    rows: list[dict[str, Any]],
    settings: Settings,
    report: IngestionReport,
) -> None:
    ids = {row["listing_id"] for row in rows}
    if report.postgres_canonical_rows != len(rows):
        raise ValueError("PostgreSQL canonical count differs from workbook count")
    with Session(engine) as session:
        stored = {
            car.listing_id: car
            for car in session.execute(
                select(Car).where(Car.listing_id.in_(ids))
            ).scalars()
        }
    if len(stored) != len(rows):
        raise ValueError("PostgreSQL listing IDs are missing")
    for row in rows:
        car = stored[row["listing_id"]]
        for field in REQUIRED:
            if getattr(car, field) != row[field]:
                raise ValueError(f"PostgreSQL {field} differs from source workbook")
        facts = extract_vehicle_facts(row["title"], row["description"])
        for field, expected in facts.items():
            if getattr(car, field) != expected:
                raise ValueError(f"PostgreSQL {field} differs from source extraction")
    store = VectorStore(settings)
    points = store.all_points()
    point_ids = [str(point.id) for point in points]
    payload_ids = [
        point.payload.get("listing_id") if point.payload else None for point in points
    ]
    if len(point_ids) != len(set(point_ids)) or len(payload_ids) != len(
        set(payload_ids)
    ):
        raise ValueError("Duplicate Qdrant point or listing IDs")
    if set(payload_ids) != ids:
        raise ValueError(
            "Qdrant listing IDs differ from canonical PostgreSQL listing IDs"
        )
    if any(str(point.id) != point_id(point.payload["listing_id"]) for point in points):
        raise ValueError(
            "Qdrant point IDs are not deterministic IDs for their listings"
        )
    if report.qdrant_vector_count != len(rows):
        raise ValueError("Qdrant vector count differs from workbook count")
    with engine.connect() as connection:
        dimensions = set(
            connection.scalars(
                select(CarEmbeddingState.dimensions).where(
                    CarEmbeddingState.listing_id.in_(ids)
                )
            )
        )
    if len(dimensions) != 1:
        raise ValueError("Missing or inconsistent recorded embedding dimensions")
    store.ensure_collection(dimensions.pop())


def run(path: Path, postgres_only: bool = False) -> IngestionReport:
    rows = read_inventory(path)
    report = IngestionReport(source_rows=len(rows))
    settings = load_settings(require_vector=not postgres_only)
    engine = make_engine(settings)
    try:
        ensure_schema(engine)
        upsert_postgres(engine, rows, report)
        if not postgres_only:
            index_vectors(engine, rows, settings, report)
            verify(engine, rows, settings, report)
        return report
    finally:
        engine.dispose()


def safe_failure(exc: Exception) -> str:
    """Classify common connection errors without echoing third-party details."""
    message = str(exc).lower()
    if "could not translate host" in message or "getaddrinfo" in message:
        return (
            "PostgreSQL host did not resolve; check Supabase host and connection mode"
        )
    if "password authentication failed" in message:
        return "PostgreSQL authentication failed; check user and password"
    if "connection refused" in message or "timed out" in message:
        return "PostgreSQL connection was refused or timed out"
    status_code = getattr(exc, "status_code", None)
    if status_code in (401, 403):
        return "Qdrant denied access; check collection name and API key"
    return f"{type(exc).__name__}; credentials and request data withheld"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument(
        "--postgres-only",
        action="store_true",
        help="Commit factual rows without vector indexing",
    )
    args = parser.parse_args()
    try:
        report = run(args.workbook, args.postgres_only)
    except (ValueError, EmbeddingError) as exc:
        print(f"Ingestion failed: {exc}", file=sys.stderr)
        return 1
    except Exception as exc:  # noqa: BLE001 - sanitize third-party errors at CLI boundary
        print(f"Ingestion failed: {safe_failure(exc)}", file=sys.stderr)
        return 1
    for label, value in asdict(report).items():
        print(f"{label.replace('_', ' ').title()}: {value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
