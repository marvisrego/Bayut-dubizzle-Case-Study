"""Build the canonical inventory with exact, content-based deduplication.

Run from the project root: python scripts/prepare_inventory.py
Requires openpyxl. The original workbook is read-only.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from inventory_text import has_noise, make_embedding_text
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = ROOT / "Cars_data.xlsx"
OUTPUT_DIR = ROOT / "data" / "processed"
SHEETS = (("raw dataset", "raw"), ("cleaned dataset", "cleaned"))
SOURCE_FIELDS = ("year", "make", "model", "trim", "title", "description", "photo_url")
MAKE_NAMES = {
    "mercedes-benz": "Mercedes-Benz",
    "mercedes benz": "Mercedes-Benz",
    "land rover": "Land Rover",
    "mini": "Mini",
    "bmw": "BMW",
    "gwm": "GWM",
    "byd": "BYD",
    "jac": "JAC",
    "mclaren": "McLaren",
    "rolls-royce": "Rolls-Royce",
    "mitsubshi": "Mitsubishi",
    "mitubshi": "Mitsubishi",
}
TOKEN_CASE = {
    "amg": "AMG",
    "bmw": "BMW",
    "gcc": "GCC",
    "gt": "GT",
    "gtc": "GTC",
    "gti": "GTI",
    "gts": "GTS",
    "hse": "HSE",
    "jcw": "JCW",
    "lx": "LX",
    "rs": "RS",
    "suv": "SUV",
    "tfsi": "TFSI",
    "v6": "V6",
    "v8": "V8",
    "4wd": "4WD",
    "4matic": "4MATIC",
    "xdrive": "xDrive",
    "xlt": "XLT",
    "sxt": "SXT",
    "glx": "GLX",
    "gxr": "GXR",
    "sv": "SV",
    "sl": "SL",
    "se": "SE",
    "le": "LE",
    "sr": "SR",
    "rst": "RST",
    "gl": "GL",
    "lc": "LC",
    "r/t": "R/T",
}
INVENTORY_COLUMNS = (
    "listing_id",
    "exact_fingerprint",
    "year",
    "make",
    "model",
    "trim",
    "title",
    "description",
    "photo_url",
    "source",
    "source_row",
    "source_listing_id",
    "source_refs",
    "year_original",
    "make_original",
    "model_original",
    "trim_original",
    "title_original",
    "description_original",
    "title_clean",
    "description_clean",
    "search_text",
    "embedding_text",
)
MERGE_LOG_COLUMNS = (
    "canonical_listing_id",
    "exact_fingerprint",
    "primary_source",
    "primary_source_row",
    "secondary_source",
    "secondary_source_row",
    "reason",
    "primary_source_record_json",
    "secondary_source_record_json",
)
CONFLICT_COLUMNS = (
    "canonical_listing_id",
    "field",
    "primary_value",
    "secondary_value",
    "resolution",
)
REVIEW_COLUMNS = (
    "year",
    "make",
    "model",
    "canonical_listing_ids",
    "distinct_fingerprints",
    "source_refs",
    "decision",
)
BR_RE = re.compile(r"<\s*br\s*/?\s*>|<\s*br\s*$", re.IGNORECASE)
TAG_RE = re.compile(r"</?[A-Za-z][^>]*>")
WHITESPACE_RE = re.compile(r"\s+")


def normalize_year(value: object) -> int | None:
    if value is None or not str(value).strip():
        return None
    try:
        number = Decimal(str(value).strip().replace(",", ""))
    except InvalidOperation:
        return None
    if number == number.to_integral_value() and 1886 <= number <= 2100:
        return int(number)
    return None


def normalize_make(value: object) -> str | None:
    if value is None or not str(value).strip():
        return None
    compact = " ".join(str(value).split())
    return MAKE_NAMES.get(compact.casefold(), compact.title())


def normalize_category(value: object) -> str | None:
    if value is None or not str(value).strip():
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    compact = " ".join(str(value).split())
    parts = re.split(r"(\s+|(?<=\w)-(?=\w))", compact)
    return "".join(
        TOKEN_CASE.get(part.casefold(), part.title())
        if part and not part.isspace() and part != "-"
        else part
        for part in parts
    )


def clean_text(value: object, *, multiline: bool) -> str | None:
    if value is None or not str(value).strip():
        return None
    result = BR_RE.sub("\n" if multiline else " ", html.unescape(str(value)))
    result = TAG_RE.sub(" ", result)
    if multiline:
        result = re.sub(r"[ \t\f\v]+", " ", result)
        result = re.sub(r" *\n *", "\n", result)
        result = re.sub(r"\n{3,}", "\n\n", result)
    else:
        result = " ".join(result.split())
    return result.strip() or None


def fingerprint_fields(record: dict) -> tuple:
    def compact(value: object) -> str:
        return (
            WHITESPACE_RE.sub(" ", unicodedata.normalize("NFKC", str(value or "")))
            .strip()
            .casefold()
        )

    return (
        record["year"],
        compact(record["make"]),
        compact(record["model"]),
        compact(record["trim"]),
        compact(record["title_clean"]),
        compact(record["description_clean"]),
    )


def exact_fingerprint(record: dict) -> str:
    payload = json.dumps(
        fingerprint_fields(record), ensure_ascii=False, separators=(",", ":")
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def read_sources(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        for sheet_name, source in SHEETS:
            sheet = workbook[sheet_name]
            rows = sheet.values
            headers = next(rows)
            positions = {
                str(name).strip().casefold(): i
                for i, name in enumerate(headers)
                if name is not None
            }
            missing = set(SOURCE_FIELDS) - positions.keys()
            if missing or (source == "cleaned" and "listing_id" not in positions):
                raise ValueError(f"{sheet_name}: missing columns {sorted(missing)}")
            seen_source_ids = set()
            for excel_row, cells in enumerate(rows, start=2):
                original = {
                    field: cells[positions[field]]
                    if positions[field] < len(cells)
                    else None
                    for field in SOURCE_FIELDS
                }
                if all(
                    value is None or str(value).strip() == ""
                    for value in original.values()
                ):
                    continue
                source_id = None
                if source == "cleaned":
                    value = (
                        cells[positions["listing_id"]]
                        if positions["listing_id"] < len(cells)
                        else None
                    )
                    if value is None or int(value) in seen_source_ids:
                        raise ValueError(
                            "Invalid or duplicate source ID at "
                            f"{sheet_name}:{excel_row}"
                        )
                    source_id = int(value)
                    seen_source_ids.add(source_id)
                record = {
                    "source": source,
                    "source_row": excel_row,
                    "source_listing_id": source_id,
                    "original": original,
                    "year": normalize_year(original["year"]),
                    "make": normalize_make(original["make"]),
                    "model": normalize_category(original["model"]),
                    "trim": normalize_category(original["trim"]),
                    "title_clean": clean_text(original["title"], multiline=False),
                    "description_clean": clean_text(
                        original["description"], multiline=True
                    ),
                }
                record["exact_fingerprint"] = exact_fingerprint(record)
                record["input_order"] = len(records)
                records.append(record)
    finally:
        workbook.close()
    return records


def source_ref(record: dict) -> dict:
    return {
        "source": record["source"],
        "source_row": record["source_row"],
        "source_listing_id": record["source_listing_id"],
    }


def source_payload(record: dict) -> str:
    return json.dumps(
        {**source_ref(record), **record["original"]}, ensure_ascii=False, default=str
    )


def make_inventory_row(primary: dict, members: list[dict], listing_id: str) -> dict:
    original = primary["original"]
    row = {
        "listing_id": listing_id,
        "exact_fingerprint": primary["exact_fingerprint"],
        "year": primary["year"],
        "make": primary["make"],
        "model": primary["model"],
        "trim": primary["trim"],
        "title": primary["title_clean"],
        "description": primary["description_clean"],
        "photo_url": original["photo_url"],
        "source": "merged" if len(members) > 1 else primary["source"],
        "source_row": primary["source_row"],
        "source_listing_id": primary["source_listing_id"],
        "source_refs": json.dumps(
            [source_ref(member) for member in members], ensure_ascii=False
        ),
        "title_clean": primary["title_clean"],
        "description_clean": primary["description_clean"],
    }
    for field in ("year", "make", "model", "trim", "title", "description"):
        row[f"{field}_original"] = original[field]
    description = row["description_clean"] or ""
    if description in (".", row["title_clean"]):
        description = ""
    row["search_text"] = ". ".join(
        part
        for part in (
            " ".join(
                str(row[field])
                for field in ("year", "make", "model", "trim")
                if row[field] is not None
            ),
            row["title_clean"],
            description,
        )
        if part
    )
    row["embedding_text"] = make_embedding_text(row)
    return row


def build(
    records: list[dict],
) -> tuple[list[dict], list[dict], list[dict], list[dict], int]:
    """Group exact duplicates by hash in expected linear time."""
    by_fingerprint = defaultdict(list)
    for record in records:
        by_fingerprint[record["exact_fingerprint"]].append(record)
    ordered_groups = sorted(
        by_fingerprint.values(),
        key=lambda group: min(record["input_order"] for record in group),
    )
    inventory, merge_log, conflicts = [], [], []
    duplicate_groups = 0
    for group in ordered_groups:
        if len({fingerprint_fields(record) for record in group}) != 1:
            raise ValueError("SHA-256 fingerprint collision detected")
        if len(group) > 1:
            duplicate_groups += 1
        primary = min(
            group,
            key=lambda record: (record["source"] != "cleaned", record["input_order"]),
        )
        members = [primary, *(record for record in group if record is not primary)]
        listing_id = f"CAR_{len(inventory) + 1:04d}"
        row = make_inventory_row(primary, members, listing_id)
        inventory.append(row)
        for secondary in members[1:]:
            merge_log.append(
                {
                    "canonical_listing_id": listing_id,
                    "exact_fingerprint": primary["exact_fingerprint"],
                    "primary_source": primary["source"],
                    "primary_source_row": primary["source_row"],
                    "secondary_source": secondary["source"],
                    "secondary_source_row": secondary["source_row"],
                    "reason": (
                        "Identical normalized year/make/model/trim/title/description"
                    ),
                    "primary_source_record_json": source_payload(primary),
                    "secondary_source_record_json": source_payload(secondary),
                }
            )
            if primary["original"]["photo_url"] != secondary["original"]["photo_url"]:
                conflicts.append(
                    {
                        "canonical_listing_id": listing_id,
                        "field": "photo_url",
                        "primary_value": primary["original"]["photo_url"],
                        "secondary_value": secondary["original"]["photo_url"],
                        "resolution": (
                            "Primary photo retained; secondary photo preserved in "
                            "merge log. Photos were not used for deduplication."
                        ),
                    }
                )

    # Retain distinct fingerprints in each blocking key for manual review.
    by_block = defaultdict(list)
    for row in inventory:
        block = tuple(row[field] for field in ("year", "make", "model"))
        if all(value is not None for value in block):
            by_block[block].append(row)
    review = []
    for (year, make, model), group in sorted(
        by_block.items(), key=lambda item: tuple(str(v) for v in item[0])
    ):
        if len(group) < 2:
            continue
        review.append(
            {
                "year": year,
                "make": make,
                "model": model,
                "canonical_listing_ids": json.dumps(
                    [row["listing_id"] for row in group]
                ),
                "distinct_fingerprints": len(group),
                "source_refs": json.dumps(
                    [json.loads(row["source_refs"]) for row in group],
                    ensure_ascii=False,
                ),
                "decision": (
                    "Kept separate: content fingerprints differ; manual review "
                    "if needed"
                ),
            }
        )
    validate(records, inventory, merge_log)
    return inventory, merge_log, conflicts, review, duplicate_groups


def validate(records: list[dict], inventory: list[dict], merge_log: list[dict]) -> None:
    source_refs = Counter(
        (record["source"], record["source_row"]) for record in records
    )
    output_refs = Counter(
        (ref["source"], ref["source_row"])
        for row in inventory
        for ref in json.loads(row["source_refs"])
    )
    if source_refs != output_refs:
        raise ValueError("A source row was lost or represented more than once")
    ids = [row["listing_id"] for row in inventory]
    fingerprints = [row["exact_fingerprint"] for row in inventory]
    if len(ids) != len(set(ids)) or len(fingerprints) != len(set(fingerprints)):
        raise ValueError("Duplicate canonical ID or exact fingerprint remains")
    if len(merge_log) != len(records) - len(inventory):
        raise ValueError("Merge log does not account for all merged source rows")
    by_ref = {(record["source"], record["source_row"]): record for record in records}
    for row in inventory:
        primary_ref = json.loads(row["source_refs"])[0]
        original = by_ref[(primary_ref["source"], primary_ref["source_row"])][
            "original"
        ]
        if (
            row["title_original"] != original["title"]
            or row["description_original"] != original["description"]
        ):
            raise ValueError(f"Original text changed for {row['listing_id']}")
        if row["year"] is not None and not isinstance(row["year"], int):
            raise ValueError(f"Invalid normalized year for {row['listing_id']}")


def write_csv(path: Path, rows: list[dict], columns: tuple[str, ...]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(
    path: Path, datasets: tuple[tuple[str, list[dict], tuple[str, ...]], ...]
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows, columns in datasets:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(columns)
        for row in rows:
            values = tuple(row.get(column) for column in columns)
            if any(len(str(value)) > 32767 for value in values if value is not None):
                raise ValueError(f"Excel cell limit exceeded in {sheet_name}")
            sheet.append(values)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def write_audit(
    path: Path,
    records: list[dict],
    inventory: list[dict],
    merge_log: list[dict],
    conflicts: list[dict],
    review: list[dict],
    duplicate_groups: int,
) -> None:
    def safe(value: object) -> str:
        return (
            str(value if value is not None else "")
            .replace("|", "\\|")
            .replace("\n", " ")
        )

    noisy = [row["listing_id"] for row in inventory if has_noise(row["embedding_text"])]
    duplicate_fingerprints = len(inventory) - len(
        {row["exact_fingerprint"] for row in inventory}
    )
    embedding_text_rows = sum(bool(row["embedding_text"]) for row in inventory)
    lines = [
        "# Exact inventory deduplication and embedding audit",
        "",
        "Generated from both sheets of Cars_data.xlsx by scripts/prepare_inventory.py.",
        "The original workbook and original text fields are unchanged.",
        "",
        "## Summary",
        "",
        f"- Rows before deduplication: {len(records)}",
        f"- Exact duplicate groups found: {duplicate_groups}",
        f"- Rows merged: {len(merge_log)}",
        f"- Final canonical rows: {len(inventory)}",
        f"- Remaining duplicate fingerprints: {duplicate_fingerprints}",
        f"- Embedding text rows created: {embedding_text_rows}",
        f"- Rows with URL/contact patterns in embedding text: {len(noisy)}",
        f"- Distinct candidate blocks retained for review: {len(review)}",
        f"- Source photo variations logged: {len(conflicts)}",
        "",
        (
            "Fingerprint inputs are normalized year, make, model, trim, title, "
            "and description."
        ),
        "The fingerprint excludes photo URL, sheet, source row, and canonical ID.",
        (
            "Candidate blocks use year + make + model only; different "
            "fingerprints stay separate."
        ),
        "Embedding text removes contact, link, address, social, and dealer boilerplate",
        "patterns. It is separate from the preserved source and clean descriptions.",
        "",
        "## Exact merges",
        "",
        "| Canonical ID | Primary source | Merged source | Fingerprint |",
        "| --- | --- | --- | --- |",
    ]
    for item in merge_log:
        lines.append(
            "| "
            + " | ".join(
                safe(value)
                for value in (
                    item["canonical_listing_id"],
                    f"{item['primary_source']}:{item['primary_source_row']}",
                    f"{item['secondary_source']}:{item['secondary_source_row']}",
                    item["exact_fingerprint"][:16],
                )
            )
            + " |"
        )
    lines.extend(
        [
            "",
            "## Candidate blocks kept separate",
            "",
            "| Year | Make | Model | Canonical IDs |",
            "| ---: | --- | --- | --- |",
        ]
    )
    for item in review:
        lines.append(
            "| "
            + " | ".join(
                safe(item[field])
                for field in ("year", "make", "model", "canonical_listing_ids")
            )
            + " |"
        )
    lines.extend(
        [
            "",
            "## Photo variations within exact groups",
            "",
            (
                "Photos did not affect matching. Full source rows remain in the "
                "Excel merge log."
            ),
            "",
            "| Canonical ID | Field | Resolution |",
            "| --- | --- | --- |",
        ]
    )
    for item in conflicts:
        lines.append(
            "| "
            + " | ".join(
                safe(item[field])
                for field in ("canonical_listing_id", "field", "resolution")
            )
            + " |"
        )
    if noisy:
        lines.extend(["", "## Residual embedding noise", "", ", ".join(noisy)])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=SOURCE_FILE)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    records = read_sources(args.input)
    inventory, merge_log, conflicts, review, duplicate_groups = build(records)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "cars_merged.csv", inventory, INVENTORY_COLUMNS)
    write_xlsx(
        args.output_dir / "cars_merged.xlsx",
        (
            ("merged_inventory", inventory, INVENTORY_COLUMNS),
            ("merge_log", merge_log, MERGE_LOG_COLUMNS),
            ("conflicts", conflicts, CONFLICT_COLUMNS),
            ("duplicate_review", review, REVIEW_COLUMNS),
        ),
    )
    write_audit(
        args.output_dir / "merge_audit.md",
        records,
        inventory,
        merge_log,
        conflicts,
        review,
        duplicate_groups,
    )
    for label, value in (
        ("Rows before deduplication", len(records)),
        ("Duplicate groups found", duplicate_groups),
        ("Rows merged", len(merge_log)),
        ("Final row count", len(inventory)),
        (
            "Remaining duplicates",
            len(inventory) - len({row["exact_fingerprint"] for row in inventory}),
        ),
        (
            "embedding_text rows created",
            sum(bool(row["embedding_text"]) for row in inventory),
        ),
        (
            "Rows still containing URLs/contact patterns",
            sum(has_noise(row["embedding_text"]) for row in inventory),
        ),
    ):
        print(f"{label}: {value}")
    print(f"Candidate blocks kept separate: {len(review)}")
    print(f"Wrote {args.output_dir / 'cars_merged.csv'}")
    print(f"Wrote {args.output_dir / 'cars_merged.xlsx'}")
    print(f"Wrote {args.output_dir / 'merge_audit.md'}")


if __name__ == "__main__":
    main()
